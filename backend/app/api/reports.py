from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import FileResponse
from app.core.database import get_db
from sqlalchemy.orm import Session, joinedload
from app.core.security import get_current_user
from app.models.evidence import Report
from app.models.test_session import TestSession
from app.models.instrument import Instrument
from app.services.audit import log_action
from pydantic import BaseModel
from typing import Optional
import hashlib
import os
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from app.core.config import settings
import uuid

router = APIRouter(prefix="/api/reports", tags=["reports"])

UPLOAD_DIR = settings.UPLOAD_DIR
os.makedirs(UPLOAD_DIR, exist_ok=True)
if settings.REPORTS_DIR and settings.REPORTS_DIR != settings.UPLOAD_DIR:
    os.makedirs(settings.REPORTS_DIR, exist_ok=True)

class ReportGenerateBody(BaseModel):
    test_session_id: Optional[int] = None
    session_id: Optional[int] = None

def _overall_result(session: TestSession) -> str:
    overall = "PASS"
    has_obs = False
    for tc in session.test_cases:
        if tc.overall_result == "FAIL":
            return "FAIL"
        for obs in tc.observations:
            has_obs = True
            if obs.result == "FAIL":
                return "FAIL"
        if tc.overall_result == "PASS":
            has_obs = True
    if not has_obs:
        return "INCOMPLETE"
    return overall

def _serialize_report(report: Report, session: TestSession = None):
    session = session or report.session
    instrument = session.instrument if session else None
    model = instrument.model if instrument else None
    return {
        "id": report.id,
        "report_number": report.report_number,
        "session_id": report.session_id,
        "status": report.status,
        "sha256_hash": report.sha256_hash,
        "version": report.version,
        "created_at": report.created_at,
        "issue_date": report.generated_at or report.created_at,
        "result": _overall_result(session) if session else None,
        "file_path": report.file_path,
        "instrument": {
            "serial_number": instrument.serial_number if instrument else None,
            "model_name": model.model_name if model else None,
            "accuracy_class": model.accuracy_class if model else None,
        } if instrument else None,
        "session_number": session.session_number if session else None,
    }

def _load_session(db: Session, session_id: int):
    return db.query(TestSession).options(
        joinedload(TestSession.instrument).joinedload(Instrument.model),
        joinedload(TestSession.test_cases),
    ).filter(TestSession.id == session_id).first()

def _generate_pdf_and_record(db: Session, session: TestSession, current_user=None):
    overall_result = _overall_result(session)
    report_number = f"TR-{session.id}-{hashlib.md5(str(session.id).encode()).hexdigest()[:6].upper()}"
    existing = db.query(Report).filter(Report.report_number == report_number).first()
    filename = f"{report_number}_{uuid.uuid4().hex[:8]}.pdf"
    file_path = os.path.join(UPLOAD_DIR, filename)
    
    c = canvas.Canvas(file_path, pagesize=letter)
    c.setFont("Helvetica-Bold", 20)
    c.drawString(50, 750, "STANDARDIZED TEST REPORT")
    
    c.setFont("Helvetica", 12)
    c.drawString(50, 720, "DEMO RULEPACK — NOT AUTHORITATIVE")
    c.drawString(50, 705, "DEMO DATA — FOR DEMONSTRATION ONLY")
    
    c.drawString(50, 670, f"Report Number: {report_number}")
    model_name = session.instrument.model.model_name if session.instrument and session.instrument.model else "-"
    serial = session.instrument.serial_number if session.instrument else "-"
    c.drawString(50, 650, f"Instrument Model: {model_name}")
    c.drawString(50, 630, f"Serial Number: {serial}")
    c.drawString(50, 610, f"Overall Result: {overall_result}")
    c.drawString(50, 590, f"Session: {session.session_number}")
    
    c.drawString(50, 560, "Observations Summary:")
    y = 540
    for tc in session.test_cases:
        c.drawString(50, y, f"Test: {tc.test_name} [{tc.overall_result or tc.status}]")
        y -= 18
        for obs in tc.observations:
            c.drawString(70, y, f"Ref: {obs.reference_value} | Ind: {obs.indicated_value} | Err: {obs.calculated_error} | Result: {obs.result}")
            y -= 14
            if y < 50:
                c.showPage()
                y = 750
    c.save()
    
    with open(file_path, "rb") as f:
        file_hash = hashlib.sha256(f.read()).hexdigest()
        
    if existing:
        existing.file_path = file_path
        existing.sha256_hash = file_hash
        existing.version = (existing.version or 1) + 1
        existing.status = "ISSUED"
        db.commit()
        db.refresh(existing)
        return existing

    report = Report(
        session_id=session.id,
        report_number=report_number,
        version=1,
        status="ISSUED",
        file_path=file_path,
        sha256_hash=file_hash,
        generated_by=current_user.id if current_user else None,
    )
    db.add(report)
    db.commit()
    db.refresh(report)
    return report

@router.post("/")
def generate_report_from_body(body: ReportGenerateBody, db: Session = Depends(get_db), current_user = Depends(get_current_user)):
    session_id = body.test_session_id or body.session_id
    if not session_id:
        raise HTTPException(status_code=400, detail="test_session_id is required")
    return generate_report(session_id, db, current_user)

@router.post("/generate/{session_id}")
def generate_report(session_id: int, db: Session = Depends(get_db), current_user = Depends(get_current_user)):
    session = _load_session(db, session_id)
    if not session:
        raise HTTPException(status_code=404, detail="Session not found")
    report = _generate_pdf_and_record(db, session, current_user)
    log_action(db, current_user.id, current_user.email, current_user.role, "GENERATE_REPORT", "Report", report.id)
    return _serialize_report(report, session)

from sqlalchemy import case

@router.get("/")
def list_reports(db: Session = Depends(get_db)):
    reports = db.query(Report).options(
        joinedload(Report.session).joinedload(TestSession.instrument).joinedload(Instrument.model),
        joinedload(Report.session).joinedload(TestSession.test_cases),
    ).order_by(
        case((Report.session_id == 1, 0), else_=1),
        Report.created_at.desc()
    ).all()
    return {"items": [_serialize_report(r) for r in reports], "total": len(reports)}

@router.get("/{id}")
def get_report(id: int, db: Session = Depends(get_db)):
    report = db.query(Report).options(
        joinedload(Report.session).joinedload(TestSession.instrument).joinedload(Instrument.model),
        joinedload(Report.session).joinedload(TestSession.test_cases),
    ).filter(Report.id == id).first()
    if not report:
        raise HTTPException(status_code=404, detail="Report not found")
    return _serialize_report(report)

@router.get("/lookup/{identifier}")
def lookup_report(identifier: str, db: Session = Depends(get_db)):
    report = None
    if identifier.isdigit():
        report = db.query(Report).filter(Report.id == int(identifier)).first()
    if not report:
        report = db.query(Report).filter(Report.report_number.ilike(identifier.strip())).first()
    if not report:
        raise HTTPException(status_code=404, detail="Report not found")
    return _serialize_report(report)

def _generate_docx(session: TestSession, report: Report) -> str:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    
    # Title
    title = doc.add_heading("STANDARDIZED TEST REPORT", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Subtitle
    sub = doc.add_paragraph("Non-Automatic Weighing Instruments (NAWI) — OIML R-76 Evaluation")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Governance Disclaimers
    disc = doc.add_paragraph()
    disc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run1 = disc.add_run("DEMO RULEPACK — NOT AUTHORITATIVE | DEMO DATA — FOR DEMONSTRATION ONLY\n")
    run1.bold = True
    run1.font.size = Pt(9)
    run1.font.color.rgb = RGBColor(180, 83, 9)
    
    # 1. Metadata Table
    doc.add_heading("1. Report & Instrument Identification", level=2)
    meta_table = doc.add_table(rows=6, cols=2)
    meta_table.style = 'Table Grid'
    
    model_name = session.instrument.model.model_name if session.instrument and session.instrument.model else "Generic Scale"
    serial = session.instrument.serial_number if session.instrument else "N/A"
    acc_class = session.instrument.model.accuracy_class if session.instrument and session.instrument.model else "III"
    overall = _overall_result(session)
    
    metadata = [
        ("Report Identifier", report.report_number),
        ("Test Session Reference", session.session_number or f"Session #{session.id}"),
        ("Instrument Model", model_name),
        ("Serial Number", serial),
        ("Accuracy Class", f"Class {acc_class}"),
        ("Deterministic Verdict", overall),
    ]
    for i, (k, v) in enumerate(metadata):
        row = meta_table.rows[i]
        row.cells[0].paragraphs[0].add_run(k).bold = True
        row.cells[1].paragraphs[0].add_run(str(v))

    # 2. Test Cases Summary
    doc.add_heading("2. Executed Test Sequences", level=2)
    if session.test_cases:
        tc_table = doc.add_table(rows=1, cols=4)
        tc_table.style = 'Table Grid'
        hdr_cells = tc_table.rows[0].cells
        hdr_cells[0].paragraphs[0].add_run("Test Sequence").bold = True
        hdr_cells[1].paragraphs[0].add_run("Procedure Type").bold = True
        hdr_cells[2].paragraphs[0].add_run("Status").bold = True
        hdr_cells[3].paragraphs[0].add_run("Result").bold = True
        
        for tc in session.test_cases:
            row_cells = tc_table.add_row().cells
            row_cells[0].text = tc.test_name or "Test"
            row_cells[1].text = tc.test_type or "OIML R-76"
            row_cells[2].text = tc.status or "COMPLETED"
            row_cells[3].text = tc.overall_result or "PASS"

    # 3. Observations Section
    doc.add_heading("3. Physical Measurement Observations", level=2)
    has_obs = any(tc.observations for tc in session.test_cases)
    if has_obs:
        obs_table = doc.add_table(rows=1, cols=5)
        obs_table.style = 'Table Grid'
        hdr_cells = obs_table.rows[0].cells
        hdr_cells[0].paragraphs[0].add_run("Test").bold = True
        hdr_cells[1].paragraphs[0].add_run("Reference Load (kg)").bold = True
        hdr_cells[2].paragraphs[0].add_run("Indicated Value (kg)").bold = True
        hdr_cells[3].paragraphs[0].add_run("Error (kg)").bold = True
        hdr_cells[4].paragraphs[0].add_run("Verdict").bold = True
        
        for tc in session.test_cases:
            for obs in tc.observations:
                row_cells = obs_table.add_row().cells
                row_cells[0].text = tc.test_name
                row_cells[1].text = str(obs.reference_value if obs.reference_value is not None else "-")
                row_cells[2].text = str(obs.indicated_value if obs.indicated_value is not None else "-")
                row_cells[3].text = str(obs.calculated_error if obs.calculated_error is not None else "-")
                row_cells[4].text = obs.result or "PASS"

    # 4. Formal Conclusion
    doc.add_heading("4. Evaluation Conclusion", level=2)
    conclusion_p = doc.add_paragraph()
    if overall == "PASS":
        conclusion_p.add_run(
            "The Non-Automatic Weighing Instrument identified in this report SATISFIES the metrological "
            "requirements under OIML Recommendation R-76 for the tested parameters."
        ).bold = True
    else:
        conclusion_p.add_run(
            "The Non-Automatic Weighing Instrument identified in this report DOES NOT SATISFY the metrological "
            "requirements under OIML Recommendation R-76 for one or more tested parameters."
        ).bold = True

    # Integrity digest note
    doc.add_paragraph(
        f"\nCryptographic Artifact Digest: {report.sha256_hash or 'Calculated upon issuance'}\n"
        "Generated by MetriSure Platform. DEMO DATA — FOR DEMONSTRATION ONLY."
    )

    docx_path = os.path.join(UPLOAD_DIR, f"{report.report_number}.docx")
    doc.save(docx_path)
    return docx_path

@router.get("/{id}/pdf")
def download_report_pdf(id: int, db: Session = Depends(get_db)):
    report = db.query(Report).filter(Report.id == id).first()
    if not report:
        raise HTTPException(status_code=404, detail="Report not found")
    if not report.file_path or not os.path.exists(report.file_path):
        if report.session_id:
            session = _load_session(db, report.session_id)
            if session:
                _generate_pdf_and_record(db, session)
                db.refresh(report)
    if not report.file_path or not os.path.exists(report.file_path):
        raise HTTPException(status_code=404, detail="Report PDF not found on disk")
    return FileResponse(report.file_path, filename=f"{report.report_number}.pdf", media_type="application/pdf")

@router.get("/{id}/docx")
def download_report_docx(id: int, db: Session = Depends(get_db)):
    report = db.query(Report).options(
        joinedload(Report.session).joinedload(TestSession.instrument).joinedload(Instrument.model),
        joinedload(Report.session).joinedload(TestSession.test_cases),
    ).filter(Report.id == id).first()
    if not report:
        raise HTTPException(status_code=404, detail="Report not found")
    if not report.session:
        raise HTTPException(status_code=400, detail="Report has no linked session")
    
    docx_path = _generate_docx(report.session, report)
    if not os.path.exists(docx_path):
        raise HTTPException(status_code=500, detail="Failed to generate DOCX report")
    return FileResponse(
        docx_path,
        filename=f"{report.report_number}.docx",
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

@router.get("/{id}/verify")
def verify_report(id: int, db: Session = Depends(get_db)):
    report = db.query(Report).options(
        joinedload(Report.session).joinedload(TestSession.test_cases),
        joinedload(Report.session).joinedload(TestSession.instrument).joinedload(Instrument.model),
    ).filter(Report.id == id).first()
    if not report:
        raise HTTPException(status_code=404, detail="Report not found")
    if (not report.file_path or not report.sha256_hash or not os.path.exists(report.file_path)) and report.session:
        _generate_pdf_and_record(db, report.session)
        db.refresh(report)
    current_hash = report.sha256_hash
    file_ok = False
    if report.file_path and os.path.exists(report.file_path):
        with open(report.file_path, "rb") as f:
            file_ok = hashlib.sha256(f.read()).hexdigest() == current_hash
    return {
        "id": report.id,
        "verified": bool(current_hash) and file_ok,
        "report_number": report.report_number,
        "status": report.status,
        "version": report.version,
        "blockchain_anchored": False,
        "hash": report.sha256_hash,
        "hash_match": file_ok,
        "overall_result": _overall_result(report.session) if report.session else None,
        "session_id": report.session_id,
        "session_number": report.session.session_number if report.session else None,
        "instrument": {
            "serial_number": report.session.instrument.serial_number if (report.session and report.session.instrument) else None,
            "model_name": report.session.instrument.model.model_name if (report.session and report.session.instrument and report.session.instrument.model) else None,
            "accuracy_class": report.session.instrument.model.accuracy_class if (report.session and report.session.instrument and report.session.instrument.model) else None,
        } if report.session else None,
        "created_at": report.created_at,
        "issue_date": report.generated_at or report.created_at,
        "disclaimer": "DEMO DATA — FOR DEMONSTRATION ONLY",
    }
